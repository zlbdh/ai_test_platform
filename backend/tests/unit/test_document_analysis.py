# -*- coding: utf-8 -*-
from io import BytesIO
import pytest
from starlette.datastructures import UploadFile

from core.document_analysis import DocumentAnalyzer
from core.document_bundle_analysis import DocumentBundleAnalyzer
from core.document_test_designer import DocumentTestDesigner
from core.models import RequirementParseRequest
from core.requirement_parser import get_requirement_parser
from routers.requirement import analyze_requirement, generate_tests_from_requirement, parse_requirement_upload


def test_document_analyzer_detects_requirement_document():
    analyzer = DocumentAnalyzer()
    content = """
    # 用户登录模块 PRD

    ## 功能需求
    1. 用户可以使用手机号+验证码登录
    2. 用户可以使用邮箱+密码登录
    3. 登录失败 3 次后锁定账户 15 分钟

    ## 验收标准
    - 密码长度不少于 8 位，必须包含大小写字母和数字
    - 验证码有效期为 5 分钟
    - 登录接口响应时间 < 500ms
    """

    analysis = analyzer.analyze(content, "登录 PRD")

    assert analysis.document_type == "requirement_prd"
    assert analysis.document_label == "需求文档"
    assert analysis.completeness_score >= 0.6
    assert analysis.testability_score >= 0.55
    assert "用户" in analysis.extracted["actors"]
    assert any("手机号+验证码登录" in flow for flow in analysis.extracted["flows"])
    assert any("密码长度不少于 8 位" in item for item in analysis.extracted["data_constraints"])
    assert "ui_e2e" in analysis.recommended_test_types
    assert "business_flow" in analysis.recommended_test_types


def test_document_analyzer_detects_api_spec():
    analyzer = DocumentAnalyzer()
    content = """
    {
      "openapi": "3.0.0",
      "paths": {
        "/users": {
          "get": {
            "summary": "List users",
            "description": "Returns all users",
            "responses": {
              "200": {"description": "ok"},
              "400": {"description": "bad request"}
            }
          }
        }
      }
    }
    """

    analysis = analyzer.analyze(content, "openapi.json")

    assert analysis.document_type == "api_spec"
    assert "GET /users" in analysis.extracted["api_endpoints"]
    assert "GET /users -> 200" in analysis.extracted["response_statuses"]
    assert "api_rest" in analysis.recommended_test_types
    assert "contract" in analysis.recommended_test_types


def test_document_analyzer_extracts_openapi_details():
    analyzer = DocumentAnalyzer()
    content = """
    {
      "openapi": "3.0.0",
      "paths": {
        "/orders/{id}": {
          "get": {
            "parameters": [
              {"name": "id", "in": "path", "required": true},
              {"name": "includeItems", "in": "query", "required": false}
            ],
            "responses": {
              "200": {"description": "ok"},
              "404": {"description": "not found"}
            }
          }
        },
        "/orders": {
          "post": {
            "requestBody": {
              "required": true,
              "content": {
                "application/json": {
                  "schema": {"$ref": "#/components/schemas/CreateOrderRequest"}
                }
              }
            },
            "responses": {
              "201": {"description": "created"},
              "400": {"description": "bad request"}
            }
          }
        }
      },
      "components": {
        "schemas": {
          "CreateOrderRequest": {
            "type": "object",
            "required": ["amount", "userId"],
            "properties": {
              "amount": {"type": "number"},
              "userId": {"type": "string"},
              "remark": {"type": "string"}
            }
          },
          "Order": {
            "type": "object",
            "properties": {
              "id": {"type": "string"}
            }
          }
        }
      }
    }
    """

    analysis = analyzer.analyze(content, "orders-openapi.json")

    assert "GET /orders/{id} -> path:id (required)" in analysis.extracted["api_parameters"]
    assert "POST /orders -> body:amount (required)" in analysis.extracted["api_parameters"]
    assert "POST /orders -> 201" in analysis.extracted["response_statuses"]
    assert "CreateOrderRequest" in analysis.extracted["schema_entities"]
    assert "Order" in analysis.extracted["schema_entities"]


def test_document_test_designer_builds_api_spec_tests():
    analyzer = DocumentAnalyzer()
    designer = DocumentTestDesigner()
    content = """
    {
      "openapi": "3.0.0",
      "paths": {
        "/orders": {
          "get": {
            "summary": "List orders",
            "responses": {
              "200": {"description": "ok"},
              "400": {"description": "bad request"}
            }
          }
        }
      }
    }
    """

    analysis = analyzer.analyze(content, "orders-openapi.json")
    designed = designer.design(content, "orders-openapi.json", analysis, None)

    assert designed["generation_summary"]["strategy_label"] == "接口契约优先"
    assert designed["generation_summary"]["counts_by_type"]["api_rest"] >= 1
    assert any(test["source"] == "api_spec" for test in designed["tests"])
    assert any("/orders" in test["name"] for test in designed["tests"])


def test_document_test_designer_builds_richer_api_spec_tests():
    analyzer = DocumentAnalyzer()
    designer = DocumentTestDesigner()
    content = """
    {
      "openapi": "3.0.0",
      "paths": {
        "/orders": {
          "post": {
            "requestBody": {
              "required": true,
              "content": {
                "application/json": {
                  "schema": {
                    "type": "object",
                    "required": ["amount"],
                    "properties": {
                      "amount": {"type": "number"},
                      "remark": {"type": "string"}
                    }
                  }
                }
              }
            },
            "responses": {
              "201": {"description": "created"},
              "400": {"description": "bad request"}
            }
          }
        }
      }
    }
    """

    analysis = analyzer.analyze(content, "orders-openapi.json")
    designed = designer.design(content, "orders-openapi.json", analysis, None)

    assert any(test["name"] == "接口参数校验: POST /orders" for test in designed["tests"])
    assert any(test["name"] == "响应状态覆盖: POST /orders" for test in designed["tests"])
    assert any(test["name"].startswith("模型结构校验:") for test in designed["tests"])


def test_document_analyzer_extracts_database_schema_details():
    analyzer = DocumentAnalyzer()
    content = """
    CREATE TABLE users (
        id BIGINT PRIMARY KEY,
        name VARCHAR(100) NOT NULL
    );

    CREATE TABLE orders (
        id BIGINT PRIMARY KEY,
        user_id BIGINT NOT NULL REFERENCES users(id),
        order_no VARCHAR(64) NOT NULL UNIQUE,
        amount DECIMAL(10,2) NOT NULL CHECK (amount > 0),
        status VARCHAR(32) DEFAULT 'pending'
    );

    CREATE INDEX idx_orders_user_id ON orders(user_id);
    """

    analysis = analyzer.analyze(content, "数据库设计.sql")

    assert analysis.document_type == "database_schema"
    assert "orders.user_id" in analysis.extracted["database_columns"]
    assert "orders.idx_orders_user_id" in analysis.extracted["database_indexes"]
    assert "orders.user_id -> users.id" in analysis.extracted["database_relations"]
    assert any("orders.amount" in item for item in analysis.extracted["data_constraints"])


def test_document_bundle_analyzer_detects_cross_document_gaps():
    analyzer = DocumentAnalyzer()
    bundle_analyzer = DocumentBundleAnalyzer()

    primary = analyzer.analyze("""
    # 登录 PRD
    1. 用户可以登录
    2. 金额必须大于 0
    3. 失败时返回 400
    """, "登录 PRD")
    reference = analyzer.analyze("""
    ## 开发文档
    POST /api/orders
    失败时记录日志
    """, "订单开发文档")

    bundle = bundle_analyzer.analyze(primary, [{
        "title": "订单开发文档",
        "document_type": reference.document_type,
        "document_label": reference.document_label,
        "extracted": reference.extracted,
    }])

    assert bundle.coverage_score <= 0.85
    assert "requirement_prd" in bundle.involved_document_types
    assert "development_design" in bundle.involved_document_types
    assert any(finding.category == "coverage" for finding in bundle.findings)
    assert any("联合" in action or "补充" in action for action in bundle.recommended_actions)


def test_document_test_designer_builds_bundle_design_from_reference_documents():
    analyzer = DocumentAnalyzer()
    designer = DocumentTestDesigner()
    parser = get_requirement_parser()

    primary_content = """
    # 订单需求
    1. 用户可以创建订单
    2. 金额必须大于 0
    3. 失败时返回 400
    """
    reference_content = """
    {
      "openapi": "3.0.0",
      "paths": {
        "/orders": {
          "post": {
            "responses": {
              "201": {"description": "created"},
              "400": {"description": "bad request"}
            }
          }
        }
      }
    }
    """

    primary_analysis = analyzer.analyze(primary_content, "订单需求")
    primary_result = parser.parse_text(primary_content, "订单需求")
    reference_analysis = analyzer.analyze(reference_content, "订单 API")
    bundle = DocumentBundleAnalyzer().analyze(primary_analysis, [{
        "title": "订单 API",
        "content": reference_content,
        "analysis": reference_analysis,
        "document_type": reference_analysis.document_type,
        "document_label": reference_analysis.document_label,
        "extracted": reference_analysis.extracted,
    }])

    designed = designer.design_bundle(
        primary_content=primary_content,
        primary_title="订单需求",
        primary_analysis=primary_analysis,
        primary_parsed_result=primary_result,
        references=[{
            "title": "订单 API",
            "content": reference_content,
            "analysis": reference_analysis,
            "parsed_result": None,
        }],
        bundle_analysis=bundle,
    )

    assert designed["generation_summary"]["strategy_label"] == "多文档联合设计"
    assert designed["generation_summary"]["counts_by_origin"]["primary"] >= 1
    assert designed["generation_summary"]["counts_by_origin"]["reference"] >= 1
    assert designed["generation_summary"]["counts_by_origin"]["bundle"] >= 1
    assert any(test["document_role"] == "reference" for test in designed["tests"])
    assert any(test["document_role"] == "bundle" for test in designed["tests"])
    assert any(source["title"] == "订单 API" for source in designed["generation_summary"]["document_sources"])


def test_document_test_designer_builds_database_schema_integrity_tests():
    analyzer = DocumentAnalyzer()
    designer = DocumentTestDesigner()
    content = """
    CREATE TABLE users (
        id BIGINT PRIMARY KEY
    );

    CREATE TABLE orders (
        id BIGINT PRIMARY KEY,
        user_id BIGINT NOT NULL REFERENCES users(id),
        amount DECIMAL(10,2) NOT NULL CHECK (amount > 0)
    );

    CREATE INDEX idx_orders_user_id ON orders(user_id);
    """

    analysis = analyzer.analyze(content, "数据库设计.sql")
    designed = designer.design(content, "数据库设计.sql", analysis, None)

    assert designed["generation_summary"]["strategy_label"] == "数据结构优先"
    assert any(test["name"] == "外键完整性校验: orders" for test in designed["tests"])
    assert any(test["name"] == "索引策略验证: orders" for test in designed["tests"])


def test_document_test_designer_builds_development_design_tests():
    analyzer = DocumentAnalyzer()
    designer = DocumentTestDesigner()
    content = """
    # 订单模块开发文档
    ## 技术方案
    1. 创建订单后调用 POST /api/orders
    2. 支付成功后更新订单状态
    3. 失败时返回 400 并记录日志
    4. 金额必须大于 0
    """

    analysis = analyzer.analyze(content, "订单模块开发文档")
    designed = designer.design(content, "订单模块开发文档", analysis, None)

    assert designed["generation_summary"]["strategy_label"] == "设计驱动"
    assert any("金额必须大于 0" in item for item in analysis.extracted["data_constraints"])
    assert any(test["source"] == "development_design" for test in designed["tests"])
    assert any("开发约束边界验证" == test["name"] for test in designed["tests"])
    assert any(test["type"] == "api_rest" for test in designed["tests"])


@pytest.mark.asyncio
async def test_requirement_analyze_route_returns_analysis_payload():
    req = RequirementParseRequest(
        title="订单需求",
        content="""
        ## 功能需求
        1. 用户可以创建订单
        2. 订单支付失败时系统必须提示失败原因

        ## 验收标准
        - 下单响应时间 < 1000ms
        - 金额必须大于 0
        """,
    )

    result = await analyze_requirement(req)

    assert result["status"] == "success"
    assert result["analysis"]["document_type"] == "requirement_prd"
    assert "quality_score" in result["analysis"]
    assert isinstance(result["analysis"]["issues"], list)
    assert isinstance(result["analysis"]["extracted"]["business_rules"], list)


@pytest.mark.asyncio
async def test_requirement_analyze_route_returns_bundle_payload():
    req = RequirementParseRequest(
        title="订单需求",
        content="""
        ## 功能需求
        1. 用户可以创建订单
        2. 金额必须大于 0
        """,
        references=[
            {
                "title": "订单开发文档",
                "content": """
                ## 技术方案
                POST /api/orders
                失败时返回 400
                """,
            }
        ],
    )

    result = await analyze_requirement(req)

    assert result["status"] == "success"
    assert "bundle_analysis" in result
    assert len(result["references_analysis"]) == 1
    assert result["bundle_analysis"]["coverage_score"] >= 0.5
    assert "development_design" in result["bundle_analysis"]["involved_document_types"]


@pytest.mark.asyncio
async def test_requirement_generate_tests_keeps_analysis_payload():
    req = RequirementParseRequest(
        title="开发设计",
        content="""
        ## 技术方案
        GET /api/orders
        POST /api/orders
        失败时返回 400
        """,
    )

    result = await generate_tests_from_requirement(req)

    assert result["status"] == "success"
    assert "analysis" in result
    assert "generation_summary" in result
    assert result["analysis"]["document_type"] in {"development_design", "api_spec"}
    assert isinstance(result["tests"], list)
    assert result["generation_summary"]["generated_count"] == len(result["tests"])
    assert any(test["type"] == "api_rest" for test in result["tests"])


@pytest.mark.asyncio
async def test_requirement_generate_tests_keeps_bundle_analysis_payload():
    req = RequirementParseRequest(
        title="订单需求",
        content="""
        ## 功能需求
        1. 用户可以创建订单
        2. 金额必须大于 0
        """,
        references=[
            {
                "title": "订单 API",
                "content": """
                {
                  "openapi": "3.0.0",
                  "paths": {
                    "/orders": {
                      "post": {
                        "responses": {
                          "201": {"description": "created"},
                          "400": {"description": "bad request"}
                        }
                      }
                    }
                  }
                }
                """,
            }
        ],
    )

    result = await generate_tests_from_requirement(req)

    assert result["status"] == "success"
    assert "bundle_analysis" in result
    assert result["bundle_analysis"]["consistency_score"] >= 0.5
    assert len(result["references_analysis"]) == 1


@pytest.mark.asyncio
async def test_requirement_generate_tests_merges_reference_generated_tests():
    req = RequirementParseRequest(
        title="订单需求",
        content="""
        ## 功能需求
        1. 用户可以创建订单
        2. 金额必须大于 0
        3. 失败时返回 400
        """,
        references=[
            {
                "title": "订单 API",
                "content": """
                {
                  "openapi": "3.0.0",
                  "paths": {
                    "/orders": {
                      "post": {
                        "responses": {
                          "201": {"description": "created"},
                          "400": {"description": "bad request"}
                        }
                      }
                    }
                  }
                }
                """,
            }
        ],
    )

    result = await generate_tests_from_requirement(req)

    assert result["status"] == "success"
    assert result["generation_summary"]["strategy_label"] == "多文档联合设计"
    assert result["generation_summary"]["counts_by_origin"]["reference"] >= 1
    assert result["generation_summary"]["counts_by_origin"]["bundle"] >= 1
    assert any(test["document_role"] == "reference" for test in result["tests"])
    assert any(test["document_role"] == "bundle" for test in result["tests"])
    assert any(source["title"] == "订单 API" for source in result["generation_summary"]["document_sources"])


@pytest.mark.asyncio
async def test_requirement_generate_tests_supports_multiple_reference_documents():
    req = RequirementParseRequest(
        title="订单需求",
        content="""
        ## 功能需求
        1. 用户可以创建订单
        2. 金额必须大于 0
        3. 失败时返回 400
        """,
        references=[
            {
                "title": "订单 API",
                "content": """
                {
                  "openapi": "3.0.0",
                  "paths": {
                    "/orders": {
                      "post": {
                        "responses": {
                          "201": {"description": "created"},
                          "400": {"description": "bad request"}
                        }
                      }
                    }
                  }
                }
                """,
            },
            {
                "title": "订单数据库设计",
                "content": """
                CREATE TABLE orders (
                    id BIGINT PRIMARY KEY,
                    amount DECIMAL(10,2) NOT NULL CHECK (amount > 0)
                );
                """,
            },
        ],
    )

    result = await generate_tests_from_requirement(req)

    assert result["status"] == "success"
    assert len(result["references_analysis"]) == 2
    assert result["generation_summary"]["counts_by_origin"]["reference"] >= 2
    assert any(source["title"] == "订单 API" for source in result["generation_summary"]["document_sources"])
    assert any(source["title"] == "订单数据库设计" for source in result["generation_summary"]["document_sources"])
    assert any(test["document_title"] == "订单 API" for test in result["tests"])
    assert any(test["document_title"] == "订单数据库设计" for test in result["tests"])
    assert "api_spec" in result["bundle_analysis"]["involved_document_types"]
    assert "database_schema" in result["bundle_analysis"]["involved_document_types"]


@pytest.mark.asyncio
async def test_requirement_generate_tests_keeps_extended_analysis_payload():
    req = RequirementParseRequest(
        title="openapi.json",
        content="""
        {
          "openapi": "3.0.0",
          "paths": {
            "/orders": {
              "post": {
                "requestBody": {
                  "required": true,
                  "content": {
                    "application/json": {
                      "schema": {
                        "type": "object",
                        "required": ["amount"],
                        "properties": {
                          "amount": {"type": "number"}
                        }
                      }
                    }
                  }
                },
                "responses": {
                  "201": {"description": "created"},
                  "400": {"description": "bad request"}
                }
              }
            }
          }
        }
        """,
    )

    result = await generate_tests_from_requirement(req)

    assert result["analysis"]["document_type"] == "api_spec"
    assert "POST /orders -> body:amount (required)" in result["analysis"]["extracted"]["api_parameters"]
    assert "POST /orders -> 201" in result["analysis"]["extracted"]["response_statuses"]
    assert result["generation_summary"]["traceability"]["api_parameters"] >= 1


@pytest.mark.asyncio
async def test_requirement_parse_upload_supports_docx():
    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_heading("登录 PRD", level=1)
    document.add_paragraph("1. 用户可以使用手机号验证码登录")
    document.add_paragraph("2. 验证码有效期为 5 分钟")
    document.save(buffer)
    buffer.seek(0)

    upload = UploadFile(filename="login.docx", file=BytesIO(buffer.getvalue()))

    result = await parse_requirement_upload(upload, title="")

    assert result["status"] == "success"
    assert result["uploaded_filename"] == "login.docx"
    assert "用户可以使用手机号验证码登录" in result["extracted_text"]
    assert result["analysis"]["document_type"] == "requirement_prd"
