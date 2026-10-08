"""
Document text extraction

Responsibilities:
- Handle txt, md, docx, pdf, and other document inputs consistently
- Provide consistent plain text for requirement parsing and test generation
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Iterable
import re

from docx import Document
from pypdf import PdfReader


TEXT_SUFFIXES = {".txt", ".md", ".json", ".sql", ".yaml", ".yml"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".docx", ".pdf"}


def extract_text_from_path(file_path: str) -> str:
    path = Path(file_path)
    data = path.read_bytes()
    return extract_text_from_bytes(path.name, data)


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    if suffix in TEXT_SUFFIXES:
        content = _decode_text_bytes(data)
    elif suffix == ".docx":
        content = _extract_docx_text(data)
    else:
        content = _extract_pdf_text(data)

    normalized = _normalize_text(content)
    if not normalized:
        raise ValueError("No usable text could be extracted from the document")
    return normalized


def _decode_text_bytes(data: bytes) -> str:
    encodings = ("utf-8", "utf-8-sig", "gb18030")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    chunks = list(_iter_docx_blocks(document))
    return "\n".join(chunks)


def _iter_docx_blocks(document: Document) -> Iterable[str]:
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            yield text

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                yield " | ".join(cells)


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages)


def _normalize_text(content: str) -> str:
    normalized = content.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()
